from io import BytesIO
from zipfile import ZipFile
import unittest
from docx import Document
from lxml import etree
from pypdf import PdfReader
from app.document_exports import build_docx_export,build_pdf_export,build_html_export,build_structured_export
from app.math_text import render_math_text
from app.models import TaskRecord
from app.note_document import normalize_note_markdown

NOTE=r'''# 混排课程

1. Parent one
   - Child bullet
     - Grandchild bullet
       wrapped continuation
   3. Child numbered
      ```python
      if 变量:
          print("日本語 🧭")  
      ```
2. Parent two

> 引用：中文标点，完整保留。

$$
\frac{x+1}{y-1} = \sqrt{z} + \alpha
$$

Inline $E=mc^{2}$ and $\frac{a}{b}$.

| Expression | Meaning |
| --- | --- |
| `a|b` | $\frac{x}{y}$ |
'''
class NestedNoteExportTests(unittest.TestCase):
    def setUp(self):
        self.task=TaskRecord(id='nested',title='混排课程',source_type='local',created_at='2026-10-08',updated_at='2026-10-08')
    def test_semantic_blocks_keep_nested_code_and_original_math(self):
        normalized=normalize_note_markdown(self.task.title,NOTE)
        self.assertFalse(normalized.report['blocking'])
        self.assertEqual(normalize_note_markdown(self.task.title,normalized.markdown).markdown,normalized.markdown)
        self.assertIn('      if 变量:\n          print("日本語 🧭")  ',normalized.markdown)
        blocks=build_structured_export(self.task,normalized.markdown)['blocks']
        self.assertEqual([(b['kind'],b['level']) for b in blocks if b['kind'] in {'ordered','bullet'}],[('ordered',0),('bullet',1),('bullet',2),('ordered',1),('ordered',0)])
        code=next(b for b in blocks if b['kind']=='code')
        self.assertEqual(code['text'],'if 变量:\n    print("日本語 🧭")  ')
        self.assertEqual(code['level'],2)
        self.assertEqual(next(b['text'] for b in blocks if b['kind']=='math'),r'\frac{x+1}{y-1} = \sqrt{z} + \alpha')
    def test_html_retains_nested_relationships_and_readable_formulas(self):
        text=build_html_export(self.task,NOTE).content.decode()
        tree=etree.HTML(text)
        self.assertEqual(len(tree.xpath('//ol/li[@value="1"]/ul/li/ul/li')),1)
        self.assertEqual(len(tree.xpath('//ol/li[@value="1"]/ol/li[@value="3"]/pre/code')),1)
        self.assertIn('(x+1)/(y-1) = sqrt(z) + α',text)
        self.assertIn('E=mc^(2)',text)
        self.assertNotIn('$$',text)
    def test_word_and_pdf_keep_editable_lists_code_tables_formulas_and_headers(self):
        result=build_docx_export(self.task,NOTE)
        document=Document(BytesIO(result.content))
        text='\n'.join(p.text for p in document.paragraphs)
        for value in ['Parent one','Parent two','Grandchild bullet wrapped continuation','(x+1)/(y-1) = sqrt(z) + α','日本語 🧭']:
            self.assertIn(value,text)
        self.assertEqual(len(document.tables),1)
        self.assertIn(self.task.title,document.sections[0].header.paragraphs[0].text)
        with ZipFile(BytesIO(result.content)) as package: root=etree.fromstring(package.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        ids=root.xpath('//w:pPr/w:numPr/w:numId/@w:val',namespaces=ns)
        self.assertEqual(ids[0],ids[2]);self.assertNotEqual(ids[0],ids[1])
        text=''.join(p.extract_text() for p in PdfReader(BytesIO(build_pdf_export(self.task,NOTE).content)).pages)
        for value in ['Parent one','Parent two','sqrt(z)','α','日本語']: self.assertIn(value,text)
        self.assertNotIn('\ufffd',text)
    def test_grouped_math_is_bounded_and_unknown_commands_remain_visible(self):
        self.assertEqual(render_math_text(r'\frac{1}{\sqrt{x}}'),'(1)/(sqrt(x))')
        self.assertEqual(render_math_text(r'\frac{x_{i+1}}{y_i}'),'(x_(i+1))/(y_i)')
        self.assertEqual(render_math_text(r'\sum_{i=1}^{n} x_i'),'Σ_(i=1)^(n) x_i')
        source=r'\unknown{literal}'
        self.assertEqual(render_math_text(source),source)
        self.assertIn('unrecognized_math_commands_preserved_as_source',build_pdf_export(self.task,'$$'+source+'$$').warnings)
    def test_diagnostic_summary_is_opt_in_and_excludes_private_details(self):
        task=self.task.model_copy(update={'status':'success','summary_source':'text-llm','summary_diagnostics':{'path':'C:/private/source','cookie':'PRIVATE_SECRET','note_quality':{'issues':[{'code':'long_paragraph','message':'PRIVATE_BODY'}]}}})
        self.assertNotIn('## 诊断摘要',build_structured_export(task,NOTE)['markdown'])
        text=build_structured_export(task,NOTE,options={'include_diagnostics':True})['markdown']
        self.assertIn('## 诊断摘要',text);self.assertIn('long_paragraph',text)
        for secret in ['PRIVATE_SECRET','PRIVATE_BODY','C:/private']:self.assertNotIn(secret,text)
